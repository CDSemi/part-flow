"""Integration tests of the Work Order file import (Phase 15 slice 1).

Exercises ``POST /api/work-orders/import/preview``, ``POST
/api/work-orders/import`` and the two template routes against a
dedicated temporary database migrated to head (SPEC §6.1: ID, AT, CC,
CR, EV, IV, AU, RC, TK, TR, TP, AZ, QB):

- Check file writes nothing and takes no lock; Import re-validates the
  checked bytes (``X-PartFlow-Import-Check``) and creates each new Work
  Order in its own transaction through the manual write path;
- an existing number — completed history included — is reported
  ``EXISTS`` and never duplicated or changed; importing the same file
  again writes nothing, which is also the recovery after a crash;
- a Work Order with any invalid row is refused whole, and a row without
  a usable number blocks the whole import;
- concurrent writers of the same number never duplicate it, a refused
  Work Order releases its PN locks while the import continues, and the
  event loop stays free during a long import;
- the audit rows carry the signed-in User and, on the Work Order's
  ``CREATED`` row only, the intake channel; reconciliation stays clean.

The API commits real transactions, so tests isolate through unique
numbers/PNs; the module database is dropped afterwards.
"""

import csv
import datetime
import hashlib
import io
import json
import os
import threading
import time
import uuid
import zipfile
from collections.abc import Callable, Iterator, Sequence
from pathlib import Path
from typing import Any, cast

import httpx
import openpyxl
import pytest
import sqlalchemy as sa
from alembic.config import Config
from fastapi.testclient import TestClient
from sqlalchemy import Engine, create_engine
from sqlalchemy.engine import URL, make_url
from sqlalchemy.orm import Session

from alembic import command
from app import cli
from app.application import part_numbers, work_orders
from app.application.errors import InvalidInputError
from app.application.work_order_import import IMPORT_COLUMNS
from app.core.config import get_settings
from app.domain.enums import Permission
from app.infrastructure import models
from app.main import create_app
from tests.auth_harness import (
    CSRF_HEADERS,
    IdentityClient,
    admin_of,
    anonymous_client,
    client_as,
    station_device_client,
)

_BACKEND_DIR = Path(__file__).resolve().parent.parent
_TEST_DATABASE = "partflow_test_work_order_import_api"
_CSV = "text/csv"
_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_CHECK = "X-PartFlow-Import-Check"
_PREVIEW = "/api/work-orders/import/preview"
_IMPORT = "/api/work-orders/import"
_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

Row = Sequence[object]


def _alembic_config(database_url: URL) -> Config:
    config = Config(str(_BACKEND_DIR / "alembic.ini"))
    config.set_main_option("script_location", str(_BACKEND_DIR / "alembic"))
    # ConfigParser interpolation reserves "%": escape the percent-encoded URL.
    url = database_url.render_as_string(hide_password=False).replace("%", "%%")
    config.set_main_option("sqlalchemy.url", url)
    return config


@pytest.fixture(scope="module")
def api_database_url() -> Iterator[URL]:
    """Temporary database migrated to head for the API under test."""
    admin_engine = create_engine(make_url(os.environ["DATABASE_URL"]), isolation_level="AUTOCOMMIT")
    with admin_engine.connect() as connection:
        connection.execute(sa.text(f'DROP DATABASE IF EXISTS "{_TEST_DATABASE}" WITH (FORCE)'))
        connection.execute(sa.text(f'CREATE DATABASE "{_TEST_DATABASE}"'))
    url = make_url(os.environ["DATABASE_URL"]).set(database=_TEST_DATABASE)
    command.upgrade(_alembic_config(url), "head")
    yield url
    with admin_engine.connect() as connection:
        connection.execute(sa.text(f'DROP DATABASE IF EXISTS "{_TEST_DATABASE}" WITH (FORCE)'))
    admin_engine.dispose()


@pytest.fixture(scope="module")
def client(api_database_url: URL) -> Iterator[TestClient]:
    """Application client wired to the temporary database."""
    original_url = os.environ["DATABASE_URL"]
    os.environ["DATABASE_URL"] = api_database_url.render_as_string(hide_password=False)
    get_settings.cache_clear()
    try:
        with TestClient(create_app()) as test_client:
            yield test_client
    finally:
        os.environ["DATABASE_URL"] = original_url
        get_settings.cache_clear()


@pytest.fixture(scope="module")
def db_engine(api_database_url: URL) -> Iterator[Engine]:
    """Direct database access for state verification."""
    engine = create_engine(api_database_url)
    yield engine
    engine.dispose()


@pytest.fixture(scope="module")
def mwo(client: TestClient) -> IdentityClient:
    """A User whose role holds only Create and edit Work Orders."""
    return client_as(client, Permission.MANAGE_WORK_ORDERS)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _unique(prefix: str) -> str:
    return f"{prefix}-{uuid.uuid4().hex[:10].upper()}"


def _csv(*rows: Row, header: Row = IMPORT_COLUMNS) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(header)
    writer.writerows(rows)
    return buffer.getvalue().encode()


def _xlsx(*rows: Row, header: Row = IMPORT_COLUMNS) -> bytes:
    """The same rows as an Excel workbook: quantities as numbers, ISO due
    dates as date cells, text as text."""
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Work Orders"
    worksheet.append(list(header))
    for row in rows:
        cells: list[object] = []
        for index, value in enumerate(row):
            if index == 2 and isinstance(value, str) and value.isdigit():
                cells.append(int(value))
            elif index == 4 and isinstance(value, str) and value:
                cells.append(datetime.date.fromisoformat(value))
            else:
                cells.append(value if value != "" else None)
        worksheet.append(cells)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _token(body: bytes) -> str:
    return hashlib.sha256(body).hexdigest()


def _preview(caller: TestClient, body: bytes, media_type: str = _CSV) -> httpx.Response:
    response: httpx.Response = caller.post(
        _PREVIEW, content=body, headers={"Content-Type": media_type}
    )
    return response


def _commit(
    caller: TestClient, body: bytes, media_type: str = _CSV, token: str | None = None
) -> httpx.Response:
    headers = {"Content-Type": media_type, _CHECK: token if token is not None else _token(body)}
    response: httpx.Response = caller.post(_IMPORT, content=body, headers=headers)
    return response


def _ok(response: httpx.Response) -> dict[str, Any]:
    assert response.status_code == 200, response.text
    return cast(dict[str, Any], response.json())


def _checked_import(caller: TestClient, body: bytes, media_type: str = _CSV) -> dict[str, Any]:
    """Check file, then Import the same bytes with the returned token."""
    preview = _ok(_preview(caller, body, media_type))
    return _ok(_commit(caller, body, media_type, preview["check_token"]))


def _entry(report: dict[str, Any], number: str) -> dict[str, Any]:
    return next(entry for entry in report["work_orders"] if entry["work_order_number"] == number)


def _outcomes(report: dict[str, Any]) -> dict[str, str]:
    return {entry["work_order_number"]: entry["outcome"] for entry in report["work_orders"]}


def _count(engine: Engine, table: sa.FromClause) -> int:
    with engine.connect() as connection:
        return connection.execute(sa.select(sa.func.count()).select_from(table)).scalar_one()


def _write_counts(engine: Engine) -> dict[str, int]:
    """Everything an import may write — and what it never may."""
    return {
        "work_orders": _count(engine, models.WorkOrder.__table__),
        "work_order_demands": _count(engine, models.WorkOrderDemand.__table__),
        "part_numbers": _count(engine, models.PartNumber.__table__),
        "audit_events": _count(engine, models.AuditEvent.__table__),
        "quantity_flows": _count(engine, models.QuantityFlow.__table__),
        "part_movements": _count(engine, models.PartMovement.__table__),
    }


def _scalar(engine: Engine, sql: str, **params: object) -> Any:
    with engine.connect() as connection:
        return connection.execute(sa.text(sql), params).scalar()


def _rows(engine: Engine, sql: str, **params: object) -> list[dict[str, Any]]:
    with engine.connect() as connection:
        return [dict(row._mapping) for row in connection.execute(sa.text(sql), params)]


def _work_order_ids(engine: Engine, numbers: Sequence[str]) -> list[int]:
    return [
        int(value)
        for value in _rows_scalar(
            engine,
            "SELECT id FROM work_orders WHERE work_order_number = ANY(:numbers) ORDER BY id",
            numbers=list(numbers),
        )
    ]


def _rows_scalar(engine: Engine, sql: str, **params: object) -> list[Any]:
    with engine.connect() as connection:
        return list(connection.execute(sa.text(sql), params).scalars())


def _number_count(engine: Engine, number: str) -> int:
    return int(
        _scalar(engine, "SELECT count(*) FROM work_orders WHERE work_order_number = :n", n=number)
    )


class _Pause:
    """Test seam: wrap ``real`` and hold its ``on_call``-th call (counted
    over every thread) until released — after the real call by default."""

    def __init__(self, real: Callable[..., Any], *, on_call: int = 1) -> None:
        self.real = real
        self.on_call = on_call
        self.inside = threading.Event()
        self.let_finish = threading.Event()
        self._guard = threading.Lock()
        self.calls = 0

    def __call__(self, *args: Any, **kwargs: Any) -> Any:
        result = self.real(*args, **kwargs)
        with self._guard:
            self.calls += 1
            should_pause = self.calls == self.on_call
        if should_pause:
            self.inside.set()
            assert self.let_finish.wait(timeout=30), "test deadlock: never released"
        return result


def _in_thread(target: Callable[[], Any]) -> tuple[threading.Thread, dict[str, Any]]:
    result: dict[str, Any] = {}

    def run() -> None:
        try:
            result["value"] = target()
        except Exception as exc:  # noqa: BLE001 — collected for assertions
            result["error"] = exc

    thread = threading.Thread(target=run, daemon=True)
    thread.start()
    return thread, result


def _identity_headers(identity_client: IdentityClient) -> dict[str, str]:
    assert identity_client.identity is not None
    return {"Cookie": f"partflow_session={identity_client.identity.token}", **CSRF_HEADERS}


# ---------------------------------------------------------------------------
# ID — the happy path, both formats, idempotent replay
# ---------------------------------------------------------------------------


def _two_work_orders() -> tuple[list[Row], str, str]:
    first, second = _unique("WO"), _unique("WO")
    rows: list[Row] = [
        (first, _unique("pn"), "3", "J-1", "2026-10-06"),
        (first, _unique("PN"), "4", "", ""),
        (first, _unique("PN"), "5", "", "2026-11-01"),
        (second, _unique("PN"), "6", "", ""),
        (second, _unique("PN"), "7", "J-2", ""),
        (second, _unique("PN"), "8", "", "2026-12-24"),
    ]
    return rows, first, second


def test_check_then_import_then_import_again(
    client: TestClient, mwo: IdentityClient, db_engine: Engine
) -> None:
    """ID-1."""
    rows, first, second = _two_work_orders()
    body = _csv(*rows)
    before = _write_counts(db_engine)
    preview = _ok(_preview(mwo, body))
    assert _write_counts(db_engine) == before
    assert preview["dry_run"] is True
    assert preview["check_token"] == _token(body)
    assert (preview["file_format"], preview["worksheet"]) == ("CSV", None)
    assert preview["summary"] == {"will_create": 2, "existing": 0, "refused": 0}
    assert (preview["rows_read"], preview["empty_rows_ignored"]) == (6, 0)
    assert preview["commit_blocked"] is False
    assert preview["lines_without_due_date"] == 3
    entry = _entry(preview, first)
    assert entry["outcome"] == "WILL_CREATE"
    assert entry["rows"] == [2, 3, 4]
    assert entry["lines"][0] == {
        "row": 2,
        "part_number": str(rows[0][1]).upper(),
        "requested_quantity": 3,
        "due_date": "2026-10-06",
        "job_number": "J-1",
    }
    assert len(entry["new_part_numbers"]) == 3
    assert (entry["work_order_id"], entry["existing_status"], entry["differs_from_file"]) == (
        None,
        None,
        None,
    )

    result = _ok(_commit(mwo, body, token=preview["check_token"]))
    assert result["dry_run"] is False
    assert result["summary"] == {"created": 2, "existing": 0, "refused": 0}
    assert _outcomes(result) == {first: "CREATED", second: "CREATED"}
    created = _write_counts(db_engine)
    assert created["work_orders"] == before["work_orders"] + 2
    assert created["work_order_demands"] == before["work_order_demands"] + 6
    assert created["part_numbers"] == before["part_numbers"] + 6
    assert created["audit_events"] == before["audit_events"] + 2 + 6 + 6
    assert _entry(result, first)["work_order_id"] == _work_order_ids(db_engine, [first])[0]

    again = _checked_import(mwo, body)
    assert again["summary"] == {"created": 0, "existing": 2, "refused": 0}
    for number in (first, second):
        assert _entry(again, number)["differs_from_file"] is False
        assert _entry(again, number)["existing_status"] == "OPEN"
    assert _write_counts(db_engine) == created


def test_xlsx_reports_like_csv(mwo: IdentityClient, db_engine: Engine) -> None:
    """ID-2: the same data as a workbook gives the same report."""
    rows, first, second = _two_work_orders()
    from_csv = _ok(_preview(mwo, _csv(*rows)))
    workbook = _xlsx(*rows)
    from_xlsx = _ok(_preview(mwo, workbook, _XLSX))
    assert (from_xlsx["file_format"], from_xlsx["worksheet"]) == ("XLSX", "Work Orders")
    ignored = {"file_format", "worksheet", "check_token"}
    assert {k: v for k, v in from_xlsx.items() if k not in ignored} == {
        k: v for k, v in from_csv.items() if k not in ignored
    }
    result = _ok(_commit(mwo, workbook, _XLSX, from_xlsx["check_token"]))
    assert _outcomes(result) == {first: "CREATED", second: "CREATED"}
    with db_engine.connect() as connection:
        due = connection.execute(
            sa.text(
                "SELECT d.due_date FROM work_order_demands d JOIN work_orders w"
                " ON w.id = d.work_order_id WHERE w.work_order_number = :n ORDER BY d.id"
            ),
            {"n": first},
        ).scalars()
        assert list(due) == [datetime.date(2026, 10, 6), None, datetime.date(2026, 11, 1)]


def test_check_file_takes_no_lock(mwo: IdentityClient, db_engine: Engine) -> None:
    """ID-3."""
    rows, _, _ = _two_work_orders()
    before = _write_counts(db_engine)
    _ok(_preview(mwo, _csv(*rows)))
    assert _write_counts(db_engine) == before
    assert _scalar(db_engine, "SELECT count(*) FROM pg_locks WHERE locktype = 'advisory'") == 0


def _manual_work_order(client: TestClient, number: str, *lines: tuple[str, int]) -> int:
    response = admin_of(client).post(
        "/api/work-orders",
        json={
            "work_order_number": number,
            "lines": [{"part_number": pn, "requested_quantity": qty} for pn, qty in lines],
        },
    )
    assert response.status_code == 201, response.text
    return int(response.json()["id"])


def test_a_completed_work_order_is_reported_never_duplicated(
    client: TestClient, mwo: IdentityClient, db_engine: Engine
) -> None:
    """ID-4 (completion set directly: only the derived status matters here)."""
    number, pn = _unique("DONE"), _unique("PN")
    work_order_id = _manual_work_order(client, number, (pn, 2))
    with db_engine.begin() as connection:
        connection.execute(
            sa.text("UPDATE work_orders SET completed_at = now() WHERE id = :id"),
            {"id": work_order_id},
        )
    before = _write_counts(db_engine)
    result = _checked_import(mwo, _csv((number, pn, "9", "", "")))
    entry = _entry(result, number)
    assert (entry["outcome"], entry["existing_status"], entry["work_order_id"]) == (
        "EXISTS",
        "COMPLETED",
        work_order_id,
    )
    assert entry["differs_from_file"] is True
    assert _write_counts(db_engine) == before
    assert _number_count(db_engine, number) == 1


def test_an_existing_open_work_order_is_not_changed(
    client: TestClient, mwo: IdentityClient, db_engine: Engine
) -> None:
    """ID-5."""
    number, pn = _unique("OPEN"), _unique("PN")
    _manual_work_order(client, number, (pn, 2))
    before = _write_counts(db_engine)
    changed_quantity = _checked_import(mwo, _csv((number, pn, "3", "", "")))
    added_line = _checked_import(
        mwo, _csv((number, pn, "2", "", ""), (number, _unique("PN"), "1", "", ""))
    )
    same = _checked_import(mwo, _csv((number, pn.lower(), "2", "J-9", "2026-10-10")))
    for report, differs in ((changed_quantity, True), (added_line, True), (same, False)):
        entry = _entry(report, number)
        assert (entry["outcome"], entry["existing_status"]) == ("EXISTS", "OPEN")
        assert entry["differs_from_file"] is differs
    assert _write_counts(db_engine) == before


# ---------------------------------------------------------------------------
# AT — one transaction per Work Order
# ---------------------------------------------------------------------------


def test_a_bad_row_refuses_only_its_work_order(mwo: IdentityClient, db_engine: Engine) -> None:
    """AT-1."""
    good, bad = _unique("A"), _unique("B")
    first_use = _unique("NEWPN")
    before = _write_counts(db_engine)
    result = _checked_import(
        mwo,
        _csv(
            (good, _unique("PN"), "1", "", ""),
            (bad, first_use, "2", "", ""),
            (bad, _unique("PN"), "0", "", ""),
        ),
    )
    assert _outcomes(result) == {good: "CREATED", bad: "REFUSED"}
    refused = _entry(result, bad)
    assert refused["lines"] == [] and refused["new_part_numbers"] == []
    assert refused["errors"] == [
        {
            "row": 4,
            "column": "Requested Quantity",
            "message": 'Requested Quantity "0" must be a positive whole number no greater than'
            " 2,147,483,647.",
        }
    ]
    after = _write_counts(db_engine)
    assert after["work_orders"] == before["work_orders"] + 1
    assert after["part_numbers"] == before["part_numbers"] + 1
    assert after["audit_events"] == before["audit_events"] + 3
    assert (
        _scalar(
            db_engine, "SELECT count(*) FROM part_numbers WHERE part_number = :pn", pn=first_use
        )
        == 0
    )


def test_a_padded_number_refuses_its_whole_work_order(
    mwo: IdentityClient, db_engine: Engine
) -> None:
    """AT-2."""
    number = _unique("WO1")
    result = _checked_import(
        mwo, _csv((number, _unique("PN"), "1", "", ""), (f"{number} ", _unique("PN"), "1", "", ""))
    )
    assert _outcomes(result) == {number: "REFUSED"}
    assert [error["row"] for error in _entry(result, number)["errors"]] == [3]
    assert _number_count(db_engine, number) == 0


def test_a_row_without_a_number_blocks_the_import(mwo: IdentityClient, db_engine: Engine) -> None:
    """AT-3."""
    body = _csv((_unique("WO"), _unique("PN"), "1", "", ""), ("", _unique("PN"), "1", "", ""))
    preview = _ok(_preview(mwo, body))
    assert preview["commit_blocked"] is True
    assert preview["unassigned_rows"] == [
        {"row": 3, "column": "Work Order Number", "message": "Work Order Number is missing."}
    ]
    before = _write_counts(db_engine)
    refused = _commit(mwo, body, token=preview["check_token"])
    assert refused.status_code == 422
    assert refused.json() == {
        "detail": "1 row has no usable Work Order Number. Add or fix it, or delete those rows,"
        " then check the file again."
    }
    assert _write_counts(db_engine) == before


def test_a_first_use_part_number_shared_by_two_work_orders(
    mwo: IdentityClient, db_engine: Engine
) -> None:
    """AT-4."""
    shared, first, second = _unique("SHARED"), _unique("WO"), _unique("WO")
    body = _csv((first, shared, "1", "", ""), (second, shared, "2", "", ""))
    preview = _ok(_preview(mwo, body))
    assert _entry(preview, first)["new_part_numbers"] == [shared]
    assert _entry(preview, second)["new_part_numbers"] == []
    result = _ok(_commit(mwo, body, token=preview["check_token"]))
    assert _outcomes(result) == {first: "CREATED", second: "CREATED"}
    assert _entry(result, first)["new_part_numbers"] == [shared]
    assert (
        _scalar(
            db_engine,
            "SELECT count(*) FROM audit_events WHERE entity_type = 'PartNumber'"
            " AND entity_id = :pn AND event_type = 'CREATED'",
            pn=shared,
        )
        == 1
    )


def test_imported_demand_lines_and_header(client: TestClient, mwo: IdentityClient) -> None:
    """AT-5."""
    number = _unique("WO")
    pns = [_unique("PN") for _ in range(3)]
    result = _checked_import(
        mwo,
        _csv(
            (number, pns[0], "5", "JOB 7", "2026-10-20"),
            (number, pns[1], "6", "", ""),
            (number, pns[2], "7", "", "2026-09-01"),
        ),
    )
    work_order_id = _entry(result, number)["work_order_id"]
    detail = admin_of(client).get(f"/api/work-orders/{work_order_id}").json()
    assert detail["work_order_number"] == number
    assert detail["due_date"] is None
    assert detail["received_date"] == datetime.date.today().isoformat()
    assert detail["status"] == "OPEN"
    demands = detail["demands"]
    assert [demand["id"] for demand in demands] == sorted(demand["id"] for demand in demands)
    assert [demand["part_number"] for demand in demands] == pns
    assert [demand["request_type"] for demand in demands] == ["NEW"] * 3
    assert [demand["due_date"] for demand in demands] == ["2026-10-20", None, "2026-09-01"]
    assert [demand["job_numbers"] for demand in demands] == [["JOB 7"], [], []]
    assert [demand["priority_rank"] for demand in demands] == [None] * 3


# ---------------------------------------------------------------------------
# CC / CR / EV — concurrency, crash recovery, the event loop
# ---------------------------------------------------------------------------


def test_a_number_created_at_flush_time_is_reported_existing(
    client: TestClient, mwo: IdentityClient, db_engine: Engine, monkeypatch: pytest.MonkeyPatch
) -> None:
    """CC-1: the unique index decides after the pre-check passed."""
    first, second = _unique("RACE"), _unique("WO")
    body = _csv((first, _unique("PN"), "1", "", ""), (second, _unique("PN"), "1", "", ""))
    token = _ok(_preview(mwo, body))["check_token"]
    pause = _Pause(part_numbers.acquire_part_number_locks)
    monkeypatch.setattr(work_orders, "acquire_part_number_locks", pause)
    thread, result = _in_thread(lambda: _commit(mwo, body, token=token))
    try:
        assert pause.inside.wait(timeout=20)
        with Session(db_engine) as session:
            work_orders.create_work_order(
                session,
                work_order_number=first,
                lines=[{"part_number": _unique("OTHER"), "requested_quantity": 1}],
                actor_user_id=admin_of(client).user_id,
            )
    finally:
        pause.let_finish.set()
    thread.join(timeout=30)
    report = _ok(result["value"])
    assert _outcomes(report) == {first: "EXISTS", second: "CREATED"}
    assert _entry(report, first)["existing_status"] == "OPEN"
    assert _number_count(db_engine, first) == 1


def test_a_number_created_before_the_pre_check_is_reported_existing(
    client: TestClient, mwo: IdentityClient, db_engine: Engine, monkeypatch: pytest.MonkeyPatch
) -> None:
    """CC-1b."""
    first, second = _unique("RACE"), _unique("WO")
    body = _csv((first, _unique("PN"), "1", "", ""), (second, _unique("PN"), "1", "", ""))
    token = _ok(_preview(mwo, body))["check_token"]
    real = work_orders.create_work_order
    raced: list[bool] = []

    def create_after_a_competitor(session: Session, **kwargs: Any) -> Any:
        if not raced:
            raced.append(True)
            with Session(db_engine) as other:
                real(
                    other,
                    work_order_number=first,
                    lines=[{"part_number": _unique("OTHER"), "requested_quantity": 1}],
                    actor_user_id=admin_of(client).user_id,
                )
        return real(session, **kwargs)

    monkeypatch.setattr(work_orders, "create_work_order", create_after_a_competitor)
    report = _ok(_commit(mwo, body, token=token))
    assert _outcomes(report) == {first: "EXISTS", second: "CREATED"}
    assert _number_count(db_engine, first) == 1


def _receiving_station(client: TestClient) -> str:
    admin = admin_of(client)
    department = admin.post("/api/departments", json={"name": _unique("DEPT")})
    assert department.status_code == 201, department.text
    area = admin.post(
        "/api/areas", json={"department_id": department.json()["id"], "name": _unique("AREA")}
    )
    assert area.status_code == 201, area.text
    area_id = area.json()["id"]
    operation = admin.post("/api/operations", json={"area_id": area_id, "code": _unique("OP")})
    assert operation.status_code == 201, operation.text
    station = admin.post(
        "/api/scan-stations", json={"station_id": _unique("ST"), "area_id": area_id}
    )
    assert station.status_code == 201, station.text
    return str(station.json()["station_id"])


def test_a_refused_work_order_releases_its_locks_during_the_import(
    client: TestClient, mwo: IdentityClient, db_engine: Engine, monkeypatch: pytest.MonkeyPatch
) -> None:
    """CC-2 and EV-1: while the import is held inside the next Work Order,
    a Scan Station receipt of the refused Work Order's PN does not wait,
    and the event loop answers other requests."""
    station_id = _receiving_station(client)
    refused_pn, first, second = _unique("LOCKED"), _unique("WO"), _unique("WO")
    body = _csv((first, refused_pn, "1", "", ""), (second, _unique("PN"), "1", "", ""))
    token = _ok(_preview(mwo, body))["check_token"]

    real_ensure = part_numbers.ensure_part_number

    def refuse_one(session: Session, value: object, **kwargs: Any) -> Any:
        if value == refused_pn:
            raise InvalidInputError("Part Number refused by the test.")
        return real_ensure(session, value, **kwargs)

    monkeypatch.setattr(work_orders, "ensure_part_number", refuse_one)
    pause = _Pause(part_numbers.acquire_part_number_locks, on_call=2)
    monkeypatch.setattr(work_orders, "acquire_part_number_locks", pause)
    # Through the module client: one event loop for every request below.
    headers = {**_identity_headers(mwo), "Content-Type": _CSV, _CHECK: token}
    thread, result = _in_thread(lambda: client.post(_IMPORT, content=body, headers=headers))
    try:
        assert pause.inside.wait(timeout=20)
        started = time.monotonic()
        health = client.get("/api/health")
        assert time.monotonic() - started < 1
        assert health.status_code == 200, health.text
        started = time.monotonic()
        receipt = station_device_client(client).post(
            f"/api/scan-stations/{station_id}/receipts",
            json={
                "part_number": refused_pn,
                "quantity": 2,
                "request_type": "MODIFY",
                "route_mode": "FLOATING",
                "scanned_at": datetime.datetime.now(datetime.UTC).isoformat(),
                "device_event_id": str(uuid.uuid4()),
            },
        )
        assert time.monotonic() - started < 2
        assert receipt.status_code == 201, receipt.text
        assert "value" not in result  # the import is still held
    finally:
        pause.let_finish.set()
    thread.join(timeout=30)
    report = _ok(result["value"])
    assert _outcomes(report) == {first: "REFUSED", second: "CREATED"}
    assert _entry(report, first)["errors"] == [
        {"row": None, "column": None, "message": "Part Number refused by the test."}
    ]


def test_two_concurrent_imports_create_each_work_order_once(
    mwo: IdentityClient, db_engine: Engine
) -> None:
    """CC-3."""
    numbers = [_unique("CC3") for _ in range(4)]
    body = _csv(*((number, _unique("PN"), "1", "", "") for number in numbers))
    token = _ok(_preview(mwo, body))["check_token"]
    barrier = threading.Barrier(2)

    def run() -> httpx.Response:
        barrier.wait(timeout=20)
        return _commit(mwo, body, token=token)

    threads = [_in_thread(run) for _ in range(2)]
    reports = []
    for thread, result in threads:
        thread.join(timeout=60)
        reports.append(_ok(result["value"]))
    for number in numbers:
        assert sorted(_outcomes(report)[number] for report in reports) == ["CREATED", "EXISTS"]
        assert _number_count(db_engine, number) == 1


def test_an_import_interrupted_by_a_crash_completes_on_retry(
    client: TestClient, mwo: IdentityClient, db_engine: Engine, monkeypatch: pytest.MonkeyPatch
) -> None:
    """CR-1."""
    numbers = [_unique("CR") for _ in range(3)]
    body = _csv(*((number, _unique("PN"), "1", "", "") for number in numbers))
    token = _ok(_preview(mwo, body))["check_token"]
    before = _write_counts(db_engine)
    real = work_orders.create_work_order

    def crash_on_the_second(session: Session, **kwargs: Any) -> Any:
        if kwargs["work_order_number"] == numbers[1]:
            raise RuntimeError("simulated crash")
        return real(session, **kwargs)

    monkeypatch.setattr(work_orders, "create_work_order", crash_on_the_second)
    raw = TestClient(client.app, headers=_identity_headers(mwo), raise_server_exceptions=False)
    crashed = raw.post(_IMPORT, content=body, headers={"Content-Type": _CSV, _CHECK: token})
    assert crashed.status_code == 500
    assert [_number_count(db_engine, number) for number in numbers] == [1, 0, 0]
    after_crash = _write_counts(db_engine)
    assert after_crash["work_orders"] == before["work_orders"] + 1
    assert after_crash["part_numbers"] == before["part_numbers"] + 1
    assert after_crash["audit_events"] == before["audit_events"] + 3

    monkeypatch.setattr(work_orders, "create_work_order", real)
    retry = _checked_import(mwo, body)
    assert _outcomes(retry) == {numbers[0]: "EXISTS", numbers[1]: "CREATED", numbers[2]: "CREATED"}
    assert _entry(retry, numbers[0])["differs_from_file"] is False
    after = _write_counts(db_engine)
    assert after["work_orders"] == before["work_orders"] + 3
    assert after["audit_events"] == before["audit_events"] + 9


# ---------------------------------------------------------------------------
# IV — text guards never reach the database as errors
# ---------------------------------------------------------------------------


def _shared_string_workbook(number_text: str, pn: str) -> bytes:
    """A workbook whose Work Order Number is a shared string written raw."""
    workbook = openpyxl.Workbook()
    buffer = io.BytesIO()
    workbook.save(buffer)
    header = "".join(
        f'<c r="{chr(65 + index)}1" t="inlineStr"><is><t>{name}</t></is></c>'
        for index, name in enumerate(IMPORT_COLUMNS[:3])
    )
    sheet = (
        f'<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="{_NS}"><sheetData>'
        f'<row r="1">{header}</row><row r="2"><c r="A2" t="s"><v>0</v></c>'
        f'<c r="B2" t="inlineStr"><is><t>{pn}</t></is></c><c r="C2"><v>1</v></c></row>'
        "</sheetData></worksheet>"
    )
    shared = (
        f'<?xml version="1.0" encoding="UTF-8"?><sst xmlns="{_NS}" count="1" uniqueCount="1">'
        f"<si><t>{number_text}</t></si></sst>"
    )
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as source,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                content = sheet.encode()
            elif info.filename == "xl/_rels/workbook.xml.rels":
                content = content.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdShared" Type="http://schemas.openxmlformats.org/'
                    b'officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"'
                    b"/></Relationships>",
                )
            elif info.filename == "[Content_Types].xml":
                content = content.replace(
                    b"</Types>",
                    b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/'
                    b'vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>',
                )
            target.writestr(info.filename, content)
        target.writestr("xl/sharedStrings.xml", shared.encode())
    return out.getvalue()


def test_text_guards_answer_without_database_errors(mwo: IdentityClient, db_engine: Engine) -> None:
    """IV-1."""
    long_number, long_pn = "W" * 3000, "P" * 3000
    body = _csv((long_number, _unique("PN"), "1", "", ""), (_unique("WO"), long_pn, "1", "", ""))
    before = _write_counts(db_engine)
    preview = _ok(_preview(mwo, body))
    assert [entry["outcome"] for entry in preview["work_orders"]] == ["REFUSED", "REFUSED"]
    assert [entry["errors"][0]["message"] for entry in preview["work_orders"]] == [
        "Work Order Number is longer than 200 characters.",
        "Part Number is longer than 200 characters.",
    ]
    result = _ok(_commit(mwo, body, token=preview["check_token"]))
    assert result["summary"] == {"created": 0, "existing": 0, "refused": 2}
    assert _write_counts(db_engine) == before

    # openpyxl may or may not decode the OOXML escape _x0000_ in a shared
    # string; either way the answer is a report, never a database error.
    escaped = f"WO_x0000_{uuid.uuid4().hex[:8]}"
    workbook = _shared_string_workbook(escaped, _unique("PN"))
    preview = _ok(_preview(mwo, workbook, _XLSX))
    (entry,) = preview["work_orders"]
    result = _ok(_commit(mwo, workbook, _XLSX, preview["check_token"]))
    if "\x00" in entry["work_order_number"]:
        assert entry["outcome"] == "REFUSED"
        assert entry["errors"][0]["message"] == (
            "Work Order Number must not contain a NUL character."
        )
        assert result["summary"]["refused"] == 1
    else:
        # openpyxl 3.1.5 keeps the escape as literal text: an ordinary number.
        assert entry["work_order_number"] == escaped
        assert entry["outcome"] == "WILL_CREATE"
        assert _outcomes(result) == {escaped: "CREATED"}


# ---------------------------------------------------------------------------
# AU / RC — audit, production data, reconciliation
# ---------------------------------------------------------------------------


def test_audit_rows_carry_the_user_and_the_intake_channel(
    client: TestClient, mwo: IdentityClient, db_engine: Engine
) -> None:
    """AU-1, AU-2, AU-3."""
    number, pn = _unique("AU"), _unique("PN")
    hot_before = _scalar(
        db_engine, "SELECT count(*) FROM work_order_demands WHERE priority_rank IS NOT NULL"
    )
    before = _write_counts(db_engine)
    result = _checked_import(mwo, _csv((number, pn, "4", "", "")))
    work_order_id = _entry(result, number)["work_order_id"]
    demand_id = _scalar(
        db_engine, "SELECT id FROM work_order_demands WHERE work_order_id = :id", id=work_order_id
    )
    rows = _rows(
        db_engine,
        "SELECT entity_type, entity_id, actor_user_id, metadata FROM audit_events"
        " WHERE (entity_type = 'WorkOrder' AND entity_id = :wo)"
        " OR (entity_type = 'WorkOrderDemand' AND entity_id = :demand)"
        " OR (entity_type = 'PartNumber' AND entity_id = :pn) ORDER BY id",
        wo=str(work_order_id),
        demand=str(demand_id),
        pn=pn,
    )
    assert {row["entity_type"]: row["metadata"] for row in rows} == {
        "PartNumber": None,
        "WorkOrder": {"intake": {"channel": "FILE_IMPORT"}},
        "WorkOrderDemand": None,
    }
    assert {row["actor_user_id"] for row in rows} == {mwo.user_id}
    after = _write_counts(db_engine)
    assert after["quantity_flows"] == before["quantity_flows"]
    assert after["part_movements"] == before["part_movements"]
    assert (
        _scalar(
            db_engine, "SELECT count(*) FROM work_order_demands WHERE priority_rank IS NOT NULL"
        )
        == hot_before
    )

    manual_id = _manual_work_order(client, _unique("MANUAL"), (_unique("PN"), 1))
    assert _rows_scalar(
        db_engine,
        "SELECT metadata FROM audit_events WHERE entity_type = 'WorkOrder' AND entity_id = :id",
        id=str(manual_id),
    ) == [None]


def test_the_audit_trail_shows_an_imported_work_order(
    client: TestClient, mwo: IdentityClient
) -> None:
    """AU-4."""
    pn = _unique("TRAIL")
    _checked_import(mwo, _csv((_unique("WO"), pn, "4", "", "")))
    response = admin_of(client).get("/api/tracking/audit-trail", params={"part_number": pn})
    body = _ok(response)
    kinds = sorted(entry["kind"] for entry in body["entries"])
    assert kinds == ["DEMAND_CREATED", "PART_NUMBER_CREATED", "WORK_ORDER_CREATED"]


def test_reconciliation_stays_clean_after_imports(
    mwo: IdentityClient, db_engine: Engine, capsys: pytest.CaptureFixture[str]
) -> None:
    """RC-1."""
    shared = _unique("SHARED")
    rows, first, second = _two_work_orders()
    third = _unique("WO")
    result = _checked_import(
        mwo, _csv(*rows, (first, shared, "2", "", ""), (third, shared, "3", "", ""))
    )
    assert set(_outcomes(result).values()) == {"CREATED"}
    work_order_ids = set(_work_order_ids(db_engine, [first, second, third]))
    demand_ids = {
        int(value)
        for value in _rows_scalar(
            db_engine,
            "SELECT id FROM work_order_demands WHERE work_order_id = ANY(:ids)",
            ids=list(work_order_ids),
        )
    }
    part_numbers = {str(row[1]).upper() for row in rows} | {shared}
    get_settings.cache_clear()
    capsys.readouterr()
    try:
        exit_code = cli.main(
            ["reconcile", "--check", "e", "--check", "f", "--check", "i", "--check", "j"]
        )
    finally:
        get_settings.cache_clear()
    report = json.loads(capsys.readouterr().out)
    assert exit_code in (0, 1), report
    named = {("WorkOrder", wo) for wo in work_order_ids} | {
        ("WorkOrderDemand", demand) for demand in demand_ids
    }
    findings = [
        finding
        for check in report["checks"]
        for finding in check["findings"]
        if (finding["entity"]["type"], finding["entity"]["id"]) in named
        or finding["part_number"] in part_numbers
    ]
    assert findings == []


# ---------------------------------------------------------------------------
# TK / TR / TP — the check token, transport refusals, templates
# ---------------------------------------------------------------------------


def test_the_import_requires_the_checked_bytes(mwo: IdentityClient, db_engine: Engine) -> None:
    """TK-1."""
    body = _csv((_unique("WO"), _unique("PN"), "1", "", ""))
    other = _csv((_unique("WO"), _unique("PN"), "1", "", ""))
    assert _ok(_preview(mwo, body))["check_token"] == hashlib.sha256(body).hexdigest()
    before = _write_counts(db_engine)
    missing = mwo.post(_IMPORT, content=body, headers={"Content-Type": _CSV})
    assert (missing.status_code, missing.json()) == (
        422,
        {"detail": "Check the file before importing it."},
    )
    for malformed in ("abc", _token(body).upper(), _token(body) + "0"):
        response = _commit(mwo, body, token=malformed)
        assert (response.status_code, response.json()["detail"]) == (
            422,
            "Check the file before importing it.",
        )
    mismatch = _commit(mwo, body, token=_token(other))
    assert (mismatch.status_code, mismatch.json()) == (
        409,
        {"detail": "This is not the file that was checked. Check the file again before importing."},
    )
    assert _write_counts(db_engine) == before


def test_media_types(mwo: IdentityClient) -> None:
    """TR-1."""
    body = _csv((_unique("WO"), _unique("PN"), "1", "", ""))
    for media_type in ("application/vnd.ms-excel", "", "text/plain"):
        response = _preview(mwo, body, media_type)
        assert (response.status_code, response.json()) == (
            415,
            {"detail": "Choose a .csv or .xlsx file."},
        )
    assert _preview(mwo, body, "text/csv; charset=utf-8").status_code == 200
    workbook = _xlsx((_unique("WO"), _unique("PN"), "1", "", ""))
    assert _preview(mwo, workbook, f"{_XLSX}; charset=binary").status_code == 200
    unlabelled = mwo.post(_PREVIEW, content=body)
    assert unlabelled.status_code == 415


def test_oversized_bodies(mwo: IdentityClient) -> None:
    """TR-2."""
    too_large = {"detail": "The file is larger than 1 MB. Split it into smaller files."}
    big = b"x" * (1_048_576 + 1)
    declared = _preview(mwo, big)
    assert (declared.status_code, declared.json()) == (413, too_large)

    def chunks() -> Iterator[bytes]:
        for _ in range(17):
            yield b"y" * 65_536

    streamed = mwo.post(_PREVIEW, content=chunks(), headers={"Content-Type": _CSV})
    assert (streamed.status_code, streamed.json()) == (413, too_large)
    lying = mwo.post(_PREVIEW, content=big, headers={"Content-Type": _CSV, "Content-Length": "10"})
    assert (lying.status_code, lying.json()) == (413, too_large)


def _rewritten_workbook(parts: dict[str, bytes]) -> bytes:
    """A saved workbook with zip members replaced or added."""
    buffer = io.BytesIO()
    openpyxl.Workbook().save(buffer)
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as source,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for info in source.infolist():
            target.writestr(info.filename, parts.get(info.filename, source.read(info.filename)))
        for name, content in parts.items():
            if name not in source.namelist():
                target.writestr(name, content)
    return out.getvalue()


def _hidden_first_sheet() -> bytes:
    workbook = openpyxl.Workbook()
    first = workbook.active
    assert first is not None
    first.title = "Hidden"
    workbook.create_sheet("Shown")
    first.sheet_state = "hidden"
    workbook.active = 1
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


_FILE_REFUSALS: list[tuple[str, bytes, str, str]] = [
    ("F2", "Work Order Number\r\nStück\r\n".encode("cp1252"), _CSV, "not saved as CSV UTF-8"),
    ("F3", b"Work Order Number\x00\r\n", _CSV, "contains a NUL character"),
    ("F4", b'Work Order Number\r\n"WO\r\n', _CSV, "cannot be read at row 2"),
    ("F5b", _xlsx(), _CSV, "This file is an Excel workbook."),
    ("F5", b"not a workbook", _XLSX, "is not an Excel workbook (.xlsx)"),
    (
        "F6",
        _rewritten_workbook({f"extra/{index}.xml": b"" for index in range(1001)}),
        _XLSX,
        "too large when unpacked",
    ),
    ("F7", _rewritten_workbook({"xl/workbook.xml": b"<workbook"}), _XLSX, "could not be read"),
    ("F14", _hidden_first_sheet(), _XLSX, "The first worksheet (Hidden) is hidden."),
    ("F8", b"\r\n , \r\n", _CSV, "The file has no header row."),
    ("F9", b"Work Order Number,Part Number\r\nWO,PN\r\n", _CSV, "Required column missing"),
    ("F10", b"Part Number,part number\r\n", _CSV, "appears more than once"),
    ("F11", b"Part Number,Quantity,Type\r\n", _CSV, "This looks like a BOM export"),
    ("F12", _csv(*(("WO", "PN", "1", "", ""),) * 2001), _CSV, "more than 2,000 data rows"),
    ("F13", _csv(), _CSV, "The file has a header row but no data rows."),
]


@pytest.mark.parametrize(
    ("body", "media_type", "message"),
    [refusal[1:] for refusal in _FILE_REFUSALS],
    ids=[refusal[0] for refusal in _FILE_REFUSALS],
)
def test_file_level_refusals(
    mwo: IdentityClient, db_engine: Engine, body: bytes, media_type: str, message: str
) -> None:
    """TR-3: through both routes, nothing written."""
    before = _write_counts(db_engine)
    for response in (_preview(mwo, body, media_type), _commit(mwo, body, media_type)):
        assert response.status_code == 422, response.text
        assert message in response.json()["detail"]
    empty = _preview(mwo, b"", media_type)
    assert (empty.status_code, empty.json()) == (422, {"detail": "The file is empty."})
    assert _write_counts(db_engine) == before


def test_templates(mwo: IdentityClient) -> None:
    """TP-1."""
    response = mwo.get("/api/work-orders/import/template.csv")
    assert response.status_code == 200
    assert response.content == (
        b"\xef\xbb\xbfWork Order Number,Part Number,Requested Quantity,Job Number,Due Date\r\n"
    )
    assert response.headers["content-type"] == "text/csv; charset=utf-8"
    assert response.headers["content-disposition"] == (
        'attachment; filename="partflow-work-order-import.csv"'
    )
    assert response.headers["cache-control"] == "no-cache"
    workbook = mwo.get("/api/work-orders/import/template.xlsx")
    assert workbook.status_code == 200
    assert workbook.headers["content-type"] == _XLSX
    assert workbook.headers["content-disposition"] == (
        'attachment; filename="partflow-work-order-import.xlsx"'
    )
    sheet = openpyxl.load_workbook(io.BytesIO(workbook.content)).active
    assert sheet is not None
    assert [cell.value for cell in sheet[1]] == list(IMPORT_COLUMNS)


# ---------------------------------------------------------------------------
# AZ — who may import
# ---------------------------------------------------------------------------


def test_only_work_order_managers_may_import(
    client: TestClient, mwo: IdentityClient, db_engine: Engine
) -> None:
    """AZ-1: authentication is judged before the media type, size and body."""
    anonymous = anonymous_client(client)
    body = _csv((_unique("WO"), _unique("PN"), "1", "", ""))
    before = _write_counts(db_engine)
    for path in (_PREVIEW, _IMPORT):
        for content, media_type in (
            (body, _CSV),
            (b"", _CSV),
            (b"x" * (2 * 1_048_576), _CSV),
            (body, "application/pdf"),
        ):
            response = anonymous.post(path, content=content, headers={"Content-Type": media_type})
            assert response.status_code == 401, response.text
            assert response.json()["authentication_required"] is True
    for template in ("template.csv", "template.xlsx"):
        response = anonymous.get(f"/api/work-orders/import/{template}")
        assert response.status_code == 401

    ewod = client_as(client, Permission.EDIT_WORK_ORDER_DEMAND)
    for response in (
        _preview(ewod, body),
        _commit(ewod, body),
        ewod.get("/api/work-orders/import/template.csv"),
        ewod.get("/api/work-orders/import/template.xlsx"),
    ):
        assert response.status_code == 403, response.text
        assert response.json()["permission_denied"] is True
        assert response.json()["required_permissions"] == ["MANAGE_WORK_ORDERS"]

    no_csrf = TestClient(client.app, headers=_identity_headers(mwo))
    del no_csrf.headers["X-PartFlow-CSRF"]
    for path in (_PREVIEW, _IMPORT):
        response = no_csrf.post(
            path, content=body, headers={"Content-Type": _CSV, _CHECK: _token(body)}
        )
        assert response.status_code == 403, response.text
        assert response.json()["csrf_rejected"] is True
    assert _write_counts(db_engine) == before
    assert _ok(_preview(mwo, body))["summary"]["will_create"] == 1


# ---------------------------------------------------------------------------
# QB — the quantity bound on the manual path
# ---------------------------------------------------------------------------


def test_the_quantity_bound_on_the_manual_path(client: TestClient, db_engine: Engine) -> None:
    """QB-1."""
    admin = admin_of(client)
    before = _write_counts(db_engine)
    too_big = admin.post(
        "/api/work-orders",
        json={"lines": [{"part_number": _unique("PN"), "requested_quantity": 2_147_483_648}]},
    )
    assert too_big.status_code == 422, too_big.text
    assert "positive whole number" in too_big.json()["detail"]
    assert _write_counts(db_engine) == before
    created = admin.post(
        "/api/work-orders",
        json={"lines": [{"part_number": _unique("PN"), "requested_quantity": 2_147_483_647}]},
    )
    assert created.status_code == 201, created.text
    line = created.json()["demands"][0]
    edited = admin.patch(
        f"/api/work-orders/{created.json()['id']}",
        json={"line_edits": [{"id": line["id"], "requested_quantity": 2_147_483_648}]},
    )
    assert edited.status_code == 422, edited.text
    assert "positive whole number" in edited.json()["detail"]
