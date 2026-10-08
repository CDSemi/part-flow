"""Tests of the Work Order import file adapters (Phase 15 slice 1, FI-1…FI-6).

No database: CSV and XLSX bytes → ``ImportSheet`` or a file-level
refusal. XLSX files are built in memory with openpyxl; where a case
needs XML openpyxl never writes (a stale or inflated ``<dimension>``, a
cached formula value, an error cell, a gap of a million rows, a
billion-laughs part), the worksheet part is replaced in the zip.
"""

import datetime
import io
import re
import struct
import time
import tracemalloc
import zipfile
from collections.abc import Callable
from typing import Literal

import openpyxl
import pytest

from app.api.work_order_import_files import (
    MAX_XLSX_CONTENT_TYPES_BYTES,
    MAX_XLSX_SHARED_STRINGS_BYTES,
    MAX_XLSX_STYLES_BYTES,
    read_csv,
    read_xlsx,
    write_xlsx_template,
)
from app.application.errors import InvalidInputError
from app.application.work_order_import import (
    IMPORT_COLUMNS,
    MAX_IMPORT_COLUMNS,
    MAX_IMPORT_RECORDS,
    CellError,
    FormattedNumber,
    ImportFormat,
    ImportSheet,
    _analyse,
)

_TOKEN = "a" * 64
_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

F1 = "The file is empty."
F2 = (
    "The file is not saved as CSV UTF-8. In Excel use Save As → CSV UTF-8 (Comma delimited),"
    " or choose the .xlsx file instead."
)
F3 = "The file contains a NUL character, so it cannot be read as CSV. Save it again as CSV UTF-8."
F5 = (
    "The file is not an Excel workbook (.xlsx). Save it as Excel Workbook (.xlsx) without a"
    " password, or as CSV UTF-8."
)
F5B = "This file is an Excel workbook. Choose a file ending in .xlsx, or save it as CSV UTF-8."
F6 = (
    "The workbook is too large when unpacked. Copy only the import columns and rows into a"
    " new workbook."
)
F7 = (
    "The workbook could not be read. Open it in Excel, save it again as Excel Workbook"
    " (.xlsx), and check it again."
)


def _csv(data: bytes) -> ImportSheet:
    return read_csv(data, check_token=_TOKEN)


def _xlsx(data: bytes) -> ImportSheet:
    return read_xlsx(data, check_token=_TOKEN)


def _refusal(reader: Callable[[bytes], ImportSheet], data: bytes) -> str:
    with pytest.raises(InvalidInputError) as caught:
        reader(data)
    return caught.value.message


def _rows(sheet: ImportSheet) -> list[tuple[int, tuple[object, ...]]]:
    return [(record.row, record.cells) for record in sheet.records]


def _save(workbook: openpyxl.Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _rewrite(data: bytes, parts: dict[str, bytes]) -> bytes:
    """Replace (or add) zip members of a saved workbook."""
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for info in source.infolist():
            target.writestr(info.filename, parts.get(info.filename, source.read(info.filename)))
        for name, content in parts.items():
            if name not in source.namelist():
                target.writestr(name, content)
    return out.getvalue()


def xlsx_with_sheet_xml(sheet_data: str, *, dimension: str | None = None) -> bytes:
    """A one-sheet workbook whose worksheet part holds ``sheet_data`` verbatim."""
    dimension_xml = f'<dimension ref="{dimension}"/>' if dimension else ""
    sheet = (
        f'<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="{_NS}">{dimension_xml}'
        f"<sheetData>{sheet_data}</sheetData></worksheet>"
    )
    return _rewrite(_save(openpyxl.Workbook()), {"xl/worksheets/sheet1.xml": sheet.encode()})


def _text(ref: str, value: str) -> str:
    return f'<c r="{ref}" t="inlineStr"><is><t>{value}</t></is></c>'


def _bounded(
    read: Callable[[], ImportSheet], *, seconds: float = 2.0, megabytes: int = 50
) -> ImportSheet:
    """Run ``read`` timed, then again under ``tracemalloc`` for its peak
    (tracing slows Python down far beyond the time bound, so the two are
    measured apart)."""
    started = time.monotonic()
    sheet = read()
    elapsed = time.monotonic() - started
    assert elapsed < seconds, elapsed
    tracemalloc.start()
    try:
        read()
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert peak < megabytes * 1024 * 1024, peak
    return sheet


# ---------------------------------------------------------------------------
# FI-1 / FI-2 — CSV
# ---------------------------------------------------------------------------


def test_csv_rows_keep_their_spreadsheet_numbers() -> None:
    """FI-1: BOM, quoted comma, leading blank line, quoted multi-line cell."""
    data = (
        "﻿\r\n"
        "Work Order Number,Part Number,Requested Quantity,Job Number\r\n"
        'WO-1,PN-1,5,"J,1"\r\n'
        'WO-1,PN-2,6,"line one\r\nline two"\r\n'
        "WO-2,PN-3,7,\r\n"
    ).encode()
    sheet = _csv(data)
    assert sheet.file_format == ImportFormat.CSV
    assert (sheet.worksheet, sheet.check_token) == (None, _TOKEN)
    assert _rows(sheet) == [
        (2, ("Work Order Number", "Part Number", "Requested Quantity", "Job Number")),
        (3, ("WO-1", "PN-1", "5", "J,1")),
        (4, ("WO-1", "PN-2", "6", "line one\r\nline two")),
        (5, ("WO-2", "PN-3", "7", "")),
    ]


def test_csv_reads_columns_a_to_bl_only() -> None:
    """FI-1: as for XLSX (OD-S1-15), cells beyond BL are not read; a row
    holding data only beyond BL is blank."""
    beyond = "," * MAX_IMPORT_COLUMNS
    data = (
        f"Work Order Number,Part Number,Requested Quantity{beyond}not read\r\n"
        f"WO-1,PN-1,5{',' * (MAX_IMPORT_COLUMNS - 3)}last column read,not read\r\n"
        f"{beyond}not read\r\n"
        "WO-2,PN-2,6\r\n"
    ).encode()
    sheet = _csv(data)
    assert [record.row for record in sheet.records] == [1, 2, 4]
    assert [len(record.cells) for record in sheet.records] == [MAX_IMPORT_COLUMNS] * 2 + [3]
    assert sheet.records[1].cells[-1] == "last column read"
    assert all("not read" not in record.cells for record in sheet.records)


def test_csv_refusals() -> None:
    """FI-2."""
    assert _refusal(_csv, b"") == F1
    assert _refusal(_csv, "Work Order Number,Part Number\r\nWO-1,Stück\r\n".encode("cp1252")) == F2
    assert _refusal(_csv, b"Work Order Number\x00,Part Number\r\n") == F3
    assert _refusal(_csv, b'Work Order Number,Part Number\r\nWO-1,"PN\r\n') == (
        "The CSV file cannot be read at row 2: a quoted value is not closed or a value is too"
        " long. Save it again as CSV UTF-8."
    )
    assert _refusal(_csv, _save(openpyxl.Workbook())) == F5B


# ---------------------------------------------------------------------------
# FI-3 — XLSX reading
# ---------------------------------------------------------------------------


def test_xlsx_values_rows_and_first_worksheet() -> None:
    """FI-3: values, an empty row 3, numbers, dates, number formats; the
    first worksheet is read and named."""
    workbook = openpyxl.Workbook()
    first = workbook.active
    assert first is not None
    first.title = "Import"
    first.append(["Work Order Number", "Part Number", "Requested Quantity", "Due Date"])
    first.append(["WO-1", "PN-1", 5, datetime.date(2026, 10, 6)])
    first["A4"] = 7010
    first["A4"].number_format = "@"
    first["B4"] = 7010
    first["C4"] = 7010
    first["C4"].number_format = "000000"
    first["D4"] = 1.23456789012346e18
    first["A5"] = "row five"
    first["BL5"] = "last column read"
    first["BM5"] = "not read"
    workbook.create_sheet("Second")["A1"] = "ignored"
    sheet = _xlsx(_save(workbook))
    assert (sheet.file_format, sheet.worksheet) == (ImportFormat.XLSX, "Import")
    assert [record.row for record in sheet.records] == [1, 2, 4, 5]
    assert all(len(record.cells) == MAX_IMPORT_COLUMNS for record in sheet.records)
    assert sheet.records[1].cells[:4] == ("WO-1", "PN-1", 5, datetime.datetime(2026, 10, 6))
    assert sheet.records[2].cells[:4] == (
        7010,
        7010,
        FormattedNumber(7010, "000000"),
        1.23456789012346e18,
    )
    assert sheet.records[3].cells[0] == "row five"
    assert sheet.records[3].cells[-1] == "last column read"
    assert "not read" not in sheet.records[3].cells


def test_xlsx_formulas_errors_and_a_stale_dimension() -> None:
    """FI-3: a cached formula value is read, one without it is empty; an
    error cell is a CellError; a stale ``<dimension>`` drops nothing."""
    rows = [
        '<row r="1">'
        + _text("A1", "Work Order Number")
        + '<c r="D1"><f>1+1</f><v>2</v></c><c r="E1"><f>1+2</f></c></row>',
        '<row r="2"><c r="B2" t="e"><v>#N/A</v></c></row>',
        *(f'<row r="{row}">{_text(f"E{row}", f"E{row}")}</row>' for row in range(3, 11)),
    ]
    sheet = _xlsx(xlsx_with_sheet_xml("".join(rows), dimension="A1:B2"))
    assert [record.row for record in sheet.records] == list(range(1, 11))
    assert sheet.records[0].cells[:5] == ("Work Order Number", None, None, 2, None)
    assert sheet.records[1].cells[1] == CellError("#N/A")
    assert sheet.records[9].cells[4] == "E10"


@pytest.mark.parametrize("state", ["hidden", "veryHidden"])
def test_a_hidden_first_worksheet_is_refused(state: Literal["hidden", "veryHidden"]) -> None:
    """FI-3."""
    workbook = openpyxl.Workbook()
    first = workbook.active
    assert first is not None
    first.title = "Old rows"
    first["A1"] = "x"
    workbook.create_sheet("Visible")["A1"] = "y"
    first.sheet_state = state
    workbook.active = 1
    assert _refusal(_xlsx, _save(workbook)) == (
        "The first worksheet (Old rows) is hidden. PartFlow reads the first worksheet only —"
        " move the sheet with the import rows to the first position and make it visible,"
        " then check the file again."
    )


# ---------------------------------------------------------------------------
# FI-4 — XLSX refusals
# ---------------------------------------------------------------------------


def _many_entries() -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        for index in range(1001):
            archive.writestr(f"part{index}.xml", b"")
    return out.getvalue()


def test_xlsx_refusals() -> None:
    """FI-4."""
    workbook = _save(openpyxl.Workbook())
    assert _refusal(_xlsx, b"") == F1
    assert _refusal(_xlsx, b"\x13\x37" * 500) == F5
    assert _refusal(_xlsx, b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 504) == F5
    assert _refusal(_xlsx, workbook[: len(workbook) // 2]) in (F5, F7)
    assert _refusal(_xlsx, _many_entries()) == F6
    inflated = _rewrite(workbook, {"xl/media/zeros.bin": b"\x00" * (17 * 1024 * 1024)})
    assert len(inflated) < 1_048_576
    assert _refusal(_xlsx, inflated) == F6
    assert _refusal(_xlsx, _rewrite(workbook, {"xl/workbook.xml": b"<workbook"})) == F7
    assert _refusal(_xlsx, _rewrite(workbook, {"xl/workbook.xml": b"not xml at all"})) == F7


def _patch_zip_headers(data: bytes, patch: Callable[[bytearray, int, int], None]) -> bytes:
    """Apply ``patch(buffer, offset, kind)`` to every local (kind 0) and
    central (kind 1) zip header, e.g. to set values zipfile never writes."""
    buffer = bytearray(data)
    for kind, signature in enumerate((b"PK\x03\x04", b"PK\x01\x02")):
        offset = buffer.find(signature)
        while offset >= 0:
            patch(buffer, offset, kind)
            offset = buffer.find(signature, offset + 4)
    return bytes(buffer)


def _encrypted(buffer: bytearray, offset: int, kind: int) -> None:
    buffer[offset + 6 + 2 * kind] |= 0x1  # general purpose flag bit 0


def _compression(method: int) -> Callable[[bytearray, int, int], None]:
    def patch(buffer: bytearray, offset: int, kind: int) -> None:
        start = offset + 8 + 2 * kind
        buffer[start : start + 2] = struct.pack("<H", method)

    return patch


def _two_row_workbook() -> bytes:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["Work Order Number", "Part Number", "Requested Quantity"])
    worksheet.append(["WO-1", "PN-1", 5])
    return _save(workbook)


def test_foreign_zip_entries_and_bad_indexes_are_refused() -> None:
    """FI-4: an encrypted entry, a compression method zipfile cannot read
    (Deflate64, AES) and an out-of-range style or shared-string index are
    refused as unreadable (F7), never an unhandled error."""
    workbook = _two_row_workbook()
    assert _refusal(_xlsx, _patch_zip_headers(workbook, _encrypted)) == F7
    for method in (9, 99):
        assert _refusal(_xlsx, _patch_zip_headers(workbook, _compression(method))) == F7
    sheet = zipfile.ZipFile(io.BytesIO(workbook)).read("xl/worksheets/sheet1.xml")
    bad_style = sheet.replace(b'<c r="C2"', b'<c r="C2" s="999"', 1)
    assert bad_style != sheet
    assert _refusal(_xlsx, _rewrite(workbook, {"xl/worksheets/sheet1.xml": bad_style})) == F7
    assert _refusal(_xlsx, _with_shared_strings("")) == F7  # index 0 of an empty table


def test_billion_laughs_is_refused_quickly() -> None:
    """FI-4: entity amplification is refused by expat (OD-S1-7)."""
    entities = '<!ENTITY lol0 "lol">' + "".join(
        f'<!ENTITY lol{level} "{f"&lol{level - 1};" * 10}">' for level in range(1, 10)
    )
    sheet = (
        f'<?xml version="1.0"?><!DOCTYPE worksheet [{entities}]><worksheet xmlns="{_NS}">'
        f'<sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>&lol9;</t></is></c></row>'
        "</sheetData></worksheet>"
    )
    data = _rewrite(_save(openpyxl.Workbook()), {"xl/worksheets/sheet1.xml": sheet.encode()})
    started = time.monotonic()
    assert _refusal(_xlsx, data) == F7
    assert time.monotonic() - started < 2


# ---------------------------------------------------------------------------
# FI-5 — the template
# ---------------------------------------------------------------------------


def test_the_xlsx_template_round_trips() -> None:
    """FI-5."""
    workbook = openpyxl.load_workbook(io.BytesIO(write_xlsx_template(IMPORT_COLUMNS)))
    worksheet = workbook.active
    assert worksheet is not None
    assert worksheet.title == "Work Orders"
    assert [cell.value for cell in worksheet[1]] == list(IMPORT_COLUMNS)
    assert worksheet.max_row == 1
    assert worksheet.freeze_panes == "A2"
    assert worksheet["A1"].font.bold
    formats = {letter: worksheet.column_dimensions[letter].number_format for letter in "ABCDE"}
    assert formats == {"A": "@", "B": "@", "C": "0", "D": "@", "E": "yyyy-mm-dd"}
    # The template itself reads as a header-only sheet.
    sheet = _xlsx(write_xlsx_template(IMPORT_COLUMNS))
    assert _rows(sheet) == [(1, (*IMPORT_COLUMNS, *(None,) * (MAX_IMPORT_COLUMNS - 5)))]


# ---------------------------------------------------------------------------
# FI-6 — amplification bounds
# ---------------------------------------------------------------------------


def test_an_inflated_dimension_pads_nothing() -> None:
    """FI-6."""
    rows = "".join(f'<row r="{row}">{_text(f"A{row}", "x")}</row>' for row in (1, 2, 3))
    data = xlsx_with_sheet_xml(rows, dimension="A1:XFD1048576")
    sheet = _bounded(lambda: _xlsx(data))
    assert [record.row for record in sheet.records] == [1, 2, 3]


def test_a_cell_on_the_last_row_stores_no_gap_records() -> None:
    """FI-6."""
    data = xlsx_with_sheet_xml(f'<row r="1048576">{_text("A1048576", "far")}</row>')
    sheet = _bounded(lambda: _xlsx(data))
    assert [record.row for record in sheet.records] == [1_048_576]
    assert sheet.records[0].cells[0] == "far"


def test_empty_cells_at_the_last_column_are_not_read() -> None:
    """FI-6: rows padded to XFD are never 16,384 cells wide."""
    rows = "".join(
        f'<row r="{row}">{_text(f"A{row}", "x") if row % 1000 == 1 else ""}<c r="XFD{row}"/></row>'
        for row in range(1, 20_001)
    )
    data = xlsx_with_sheet_xml(rows, dimension="A1:XFD20000")
    sheet = _bounded(lambda: _xlsx(data))
    assert [record.row for record in sheet.records] == list(range(1, 20_001, 1000))
    assert {len(record.cells) for record in sheet.records} == {MAX_IMPORT_COLUMNS}


_SHARED_STRINGS_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"
)
_SHARED_STRINGS_REL = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings"
)


def _with_shared_strings(entries: str, *, part: str = "xl/sharedStrings.xml") -> bytes:
    """The two-row workbook with cell A2 read from index 0 of a shared-string
    part holding ``entries``, registered at ``part``."""
    workbook = _two_row_workbook()
    source = zipfile.ZipFile(io.BytesIO(workbook))
    rels = source.read("xl/_rels/workbook.xml.rels").replace(
        b"</Relationships>",
        f'<Relationship Id="rIdS" Type="{_SHARED_STRINGS_REL}" Target="/{part}"/>'
        "</Relationships>".encode(),
    )
    types = source.read("[Content_Types].xml").replace(
        b"</Types>",
        f'<Override PartName="/{part}" ContentType="{_SHARED_STRINGS_TYPE}"/></Types>'.encode(),
    )
    sheet, replaced = re.subn(
        rb'<c r="A2" t="inlineStr">.*?</c>',
        b'<c r="A2" t="s"><v>0</v></c>',
        source.read("xl/worksheets/sheet1.xml"),
        count=1,
    )
    assert replaced == 1
    strings = f'<?xml version="1.0" encoding="UTF-8"?><sst xmlns="{_NS}">{entries}</sst>'
    return _rewrite(
        workbook,
        {
            "xl/_rels/workbook.xml.rels": rels,
            "[Content_Types].xml": types,
            "xl/worksheets/sheet1.xml": sheet,
            part: strings.encode(),
        },
    )


def _with_cell_formats(declared_bytes: int) -> bytes:
    """The two-row workbook whose ``cellXfs`` is padded with empty ``<xf/>``
    entries to about ``declared_bytes`` of ``xl/styles.xml``."""
    workbook = _two_row_workbook()
    styles = zipfile.ZipFile(io.BytesIO(workbook)).read("xl/styles.xml").decode()
    padding = "<xf/>" * ((declared_bytes - len(styles)) // len("<xf/>"))
    padded, replaced = re.subn(r'<cellXfs count="\d+">', rf"\g<0>{padding}", styles, count=1)
    assert replaced == 1
    return _rewrite(workbook, {"xl/styles.xml": padded.encode()})


def test_a_relocated_shared_strings_part_is_read() -> None:
    """FI-3: the content types, not the part name, locate shared strings."""
    sheet = _xlsx(_with_shared_strings("<si><t>WO-S</t></si>", part="xl/strings.xml"))
    assert sheet.records[1].cells[:3] == ("WO-S", "PN-1", 5)


def test_style_and_shared_string_tables_are_bounded() -> None:
    """FI-6: a tiny workbook of empty style or shared-string entries —
    tables openpyxl builds whole at load — is read within the budget just
    below its part bound and refused (F6) above it, wherever the content
    types place the shared strings."""
    below_styles = _with_cell_formats(MAX_XLSX_STYLES_BYTES - 64)
    assert len(below_styles) < 1_048_576
    assert len(_bounded(lambda: _xlsx(below_styles)).records) == 2
    assert _refusal(_xlsx, _with_cell_formats(MAX_XLSX_STYLES_BYTES + 64)) == F6

    def strings(declared_bytes: int, part: str = "xl/sharedStrings.xml") -> bytes:
        padding = "<si/>" * (declared_bytes // len("<si/>"))
        return _with_shared_strings("<si><t>WO-S</t></si>" + padding, part=part)

    below_strings = strings(MAX_XLSX_SHARED_STRINGS_BYTES - 256)
    assert len(_bounded(lambda: _xlsx(below_strings)).records) == 2
    assert _refusal(_xlsx, strings(MAX_XLSX_SHARED_STRINGS_BYTES)) == F6
    assert _refusal(_xlsx, strings(MAX_XLSX_SHARED_STRINGS_BYTES, "xl/worksheets/s.xml")) == F6

    workbook = _two_row_workbook()
    types = zipfile.ZipFile(io.BytesIO(workbook)).read("[Content_Types].xml")
    comment = b"<!--" + b"x" * MAX_XLSX_CONTENT_TYPES_BYTES + b"-->"
    padded = types.replace(b"</Types>", comment + b"</Types>")
    assert _refusal(_xlsx, _rewrite(workbook, {"[Content_Types].xml": padded})) == F6


def test_a_wide_csv_row_costs_no_more_than_64_columns() -> None:
    """FI-6: one row of about a million commas among 2,000 short rows (just
    under 1 MB) is read as columns A–BL, so the header check stays cheap."""
    data = (
        "Work Order Number,Part Number,Requested Quantity\n"
        + "WO-1,PN-1,1"
        + "," * 1_000_000
        + "\n"
        + "WO-1,PN-2,1\n" * 1999
    ).encode()
    assert len(data) < 1_048_576
    sheet = _bounded(lambda: _csv(data))
    assert max(len(record.cells) for record in sheet.records) == MAX_IMPORT_COLUMNS
    started = time.monotonic()
    analysis = _analyse(sheet)
    assert time.monotonic() - started < 2
    assert analysis.ignored_columns == ()


def test_csv_blank_lines_are_never_stored() -> None:
    """FI-6: 1,048,576 newlines → zero records (the Application says F8)."""
    sheet = _bounded(lambda: _csv(b"\n" * 1_048_576))
    assert sheet.records == ()
    with pytest.raises(InvalidInputError, match="The file has no header row"):
        _analyse(sheet)


def test_reading_stops_after_the_row_limit() -> None:
    """FI-6: 3,000 rows → the adapter keeps MAX_IMPORT_RECORDS; F12 follows."""
    data = ("Work Order Number,Part Number,Requested Quantity\n" + "WO,PN,1\n" * 3000).encode()
    sheet = _csv(data)
    assert len(sheet.records) == MAX_IMPORT_RECORDS
    with pytest.raises(InvalidInputError, match="more than 2,000 data rows"):
        _analyse(sheet)
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    for _ in range(3000):
        worksheet.append(["WO", "PN", 1])
    assert len(_xlsx(_save(workbook)).records) == MAX_IMPORT_RECORDS
