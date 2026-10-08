"""File-format adapters of the Work Order import (Phase 15 slice 1).

Bytes → :class:`~app.application.work_order_import.ImportSheet`, or a
file-level ``InvalidInputError`` (422, nothing written). The adapters
know CSV and XLSX — encodings, the zip container, NUL bytes, row
numbering — and nothing about column names: the header contract, the
limits and every cell rule live in ``app.application.work_order_import``
(one rule source for both formats). This is the only module under
``app/`` that imports ``openpyxl``.

Both adapters iterate lazily, read columns A–BL only
(``MAX_IMPORT_COLUMNS``: cells beyond are never read, so a crafted
row cannot widen the work done per record), never store a fully blank
record (its row number is all the Application needs) and stop once
``MAX_IMPORT_RECORDS`` non-blank records are kept — enough for the
Application to refuse a file over the row limit.

XLSX is read values-only (OD-15-1/2): ``data_only=True`` reads the value
Excel stored when it saved (a formula is never evaluated; one without
a stored value reads as empty), macros are never run, the first
worksheet is read (refused when hidden) and the stored ``<dimension>``
is never trusted. The zip guard bounds the XML that is parsed — in
total, and separately for the parts openpyxl parses whole into object
tables at load (content types, styles, shared strings; the worksheet
is streamed); Python's expat (≥ 2.4.1) bounds entity amplification.
"""

import csv
import datetime
import io
import xml.etree.ElementTree
import zipfile
import zlib
from typing import Any, Final, cast

import openpyxl
from openpyxl.cell.read_only import EMPTY_CELL, ReadOnlyCell
from openpyxl.packaging.manifest import Manifest
from openpyxl.styles import Font
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from openpyxl.xml.constants import ARC_CONTENT_TYPES, ARC_STYLE, SHARED_STRINGS
from openpyxl.xml.functions import fromstring

from app.application.errors import InvalidInputError, UnsupportedMediaTypeError
from app.application.work_order_import import (
    COLUMN_DUE_DATE,
    COLUMN_JOB_NUMBER,
    COLUMN_PART_NUMBER,
    COLUMN_REQUESTED_QUANTITY,
    COLUMN_WORK_ORDER_NUMBER,
    MAX_IMPORT_COLUMNS,
    MAX_IMPORT_RECORDS,
    CellError,
    CellValue,
    FormattedNumber,
    ImportFormat,
    ImportRecord,
    ImportSheet,
    is_blank_cell,
)

#: Largest import body (OD-15-10).
MAX_IMPORT_BYTES: Final = 1_048_576
#: Zip guard: entries and total declared uncompressed size of a workbook.
MAX_XLSX_ENTRIES: Final = 1000
MAX_XLSX_UNPACKED_BYTES: Final = 16 * 1024 * 1024
#: Zip guard per part openpyxl parses whole at load: its object tables
#: grow with the entry count, not with the bytes kept (a tiny workbook
#: of empty ``<xf/>`` or ``<si/>`` entries costs seconds and hundreds of
#: MB below the total bound). Each bound keeps the worst case of its
#: part inside the FI-6 budget (< 2 s, < 50 MB) and far above a real
#: workbook's part.
MAX_XLSX_CONTENT_TYPES_BYTES: Final = 256 * 1024
MAX_XLSX_STYLES_BYTES: Final = 256 * 1024
MAX_XLSX_SHARED_STRINGS_BYTES: Final = 2 * 1024 * 1024

CSV_MEDIA_TYPE: Final = "text/csv"
XLSX_MEDIA_TYPE: Final = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

FILE_TOO_LARGE_MESSAGE: Final = "The file is larger than 1 MB. Split it into smaller files."
UNSUPPORTED_FILE_MESSAGE: Final = "Choose a .csv or .xlsx file."
EMPTY_FILE_MESSAGE: Final = "The file is empty."
NOT_UTF8_MESSAGE: Final = (
    "The file is not saved as CSV UTF-8. In Excel use Save As → CSV UTF-8 (Comma delimited),"
    " or choose the .xlsx file instead."
)
CSV_NUL_MESSAGE: Final = (
    "The file contains a NUL character, so it cannot be read as CSV. Save it again as CSV UTF-8."
)
NOT_XLSX_MESSAGE: Final = (
    "The file is not an Excel workbook (.xlsx). Save it as Excel Workbook (.xlsx) without a"
    " password, or as CSV UTF-8."
)
WORKBOOK_AS_CSV_MESSAGE: Final = (
    "This file is an Excel workbook. Choose a file ending in .xlsx, or save it as CSV UTF-8."
)
WORKBOOK_TOO_LARGE_MESSAGE: Final = (
    "The workbook is too large when unpacked. Copy only the import columns and rows into a"
    " new workbook."
)
WORKBOOK_UNREADABLE_MESSAGE: Final = (
    "The workbook could not be read. Open it in Excel, save it again as Excel Workbook"
    " (.xlsx), and check it again."
)

_UTF8_BOM: Final = b"\xef\xbb\xbf"
_ZIP_SIGNATURE: Final = b"PK\x03\x04"
#: The failures of reading a damaged or foreign workbook: the file is
#: refused (F7), nothing is written. ``RuntimeError``: an encrypted zip
#: entry; ``NotImplementedError``: a compression method zipfile cannot
#: read; ``IndexError``: a style or shared-string index out of range.
_WORKBOOK_ERRORS: Final = (
    zipfile.BadZipFile,
    InvalidFileException,
    KeyError,
    ValueError,
    TypeError,
    IndexError,
    RuntimeError,
    NotImplementedError,
    xml.etree.ElementTree.ParseError,
    OSError,
    EOFError,
    zlib.error,
)
_TEXT_FORMATS: Final = frozenset({"general", "@"})
#: The filler openpyxl puts in place of an absent cell.
_FILLER: Final[object] = EMPTY_CELL


def import_format_of(content_type: str | None) -> ImportFormat:
    """The import format named by the request's media type (415 otherwise).

    Parameters such as ``; charset=utf-8`` are allowed; the client labels
    the body by file extension, never by the browser's guess.
    """
    media_type = (content_type or "").split(";", 1)[0].strip().lower()
    if media_type == CSV_MEDIA_TYPE:
        return ImportFormat.CSV
    if media_type == XLSX_MEDIA_TYPE:
        return ImportFormat.XLSX
    raise UnsupportedMediaTypeError(UNSUPPORTED_FILE_MESSAGE)


def reject_empty(data: bytes) -> None:
    if not data:
        raise InvalidInputError(EMPTY_FILE_MESSAGE)


def read_import_file(data: bytes, file_format: ImportFormat, *, check_token: str) -> ImportSheet:
    if file_format == ImportFormat.CSV:
        return read_csv(data, check_token=check_token)
    return read_xlsx(data, check_token=check_token)


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------


def read_csv(data: bytes, *, check_token: str) -> ImportSheet:
    """UTF-8 (optional BOM), comma, RFC 4180 quoting; columns A–BL.

    The row number of a record is its 1-based index among the records
    ``csv.reader`` yields (a blank line yields one and still counts), so
    it stays the spreadsheet row after a quoted multi-line cell.
    """
    reject_empty(data)
    if data.startswith(_UTF8_BOM):
        data = data[len(_UTF8_BOM) :]
    if data.startswith(_ZIP_SIGNATURE):
        raise InvalidInputError(WORKBOOK_AS_CSV_MESSAGE)
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise InvalidInputError(NOT_UTF8_MESSAGE) from exc
    if "\x00" in text:
        raise InvalidInputError(CSV_NUL_MESSAGE)
    records: list[ImportRecord] = []
    reader = csv.reader(io.StringIO(text, newline=""), strict=True)
    row = 0
    try:
        for fields in reader:
            row += 1
            cells = fields[:MAX_IMPORT_COLUMNS]
            if cells and not all(is_blank_cell(cell) for cell in cells):
                records.append(ImportRecord(row=row, cells=tuple(cells)))
                if len(records) >= MAX_IMPORT_RECORDS:
                    break
    except csv.Error as exc:
        raise InvalidInputError(
            f"The CSV file cannot be read at row {row + 1}: a quoted value is not closed"
            " or a value is too long. Save it again as CSV UTF-8."
        ) from exc
    return ImportSheet(
        file_format=ImportFormat.CSV,
        records=tuple(records),
        check_token=check_token,
        worksheet=None,
    )


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------


class _UnreadableCellError(Exception):
    """A cell reported another row than its position (never shift silently),
    or holds a value type a worksheet never stores."""


def _cell_value(cell: object, row: int) -> CellValue:
    if not isinstance(cell, ReadOnlyCell):
        return None  # filler for an absent cell
    if cell.row != row:
        raise _UnreadableCellError
    value = cell.value
    if value is None:
        return None
    if cell.data_type == "e":
        return CellError(str(value))
    if isinstance(value, int | float) and not isinstance(value, bool):
        number_format = str(cell.number_format or "General")
        if number_format.casefold() not in _TEXT_FORMATS:
            return FormattedNumber(value, number_format)
        return value
    if isinstance(value, str | datetime.date | datetime.time | datetime.timedelta):
        return value  # datetime included
    raise _UnreadableCellError


def _shared_strings_part(archive: zipfile.ZipFile) -> str | None:
    """The shared-string part openpyxl reads: the first override of its
    content type in ``[Content_Types].xml``, resolved as openpyxl does."""
    # The stubs type ``from_tree``'s node narrower than ``fromstring``'s
    # result; at run time openpyxl passes exactly this element.
    root = cast(Any, fromstring(archive.read(ARC_CONTENT_TYPES)))
    manifest = Manifest.from_tree(root)
    part = None if manifest is None else next(manifest.findall(SHARED_STRINGS), None)
    return None if part is None else str(part.PartName)[1:]


def _check_zip(data: bytes) -> None:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise InvalidInputError(NOT_XLSX_MESSAGE) from exc
    with archive:
        entries = archive.infolist()
        if (
            len(entries) > MAX_XLSX_ENTRIES
            or sum(entry.file_size for entry in entries) > MAX_XLSX_UNPACKED_BYTES
        ):
            raise InvalidInputError(WORKBOOK_TOO_LARGE_MESSAGE)
        # zipfile resolves a duplicated name to its last entry, as here.
        sizes = {entry.filename: entry.file_size for entry in entries}
        if (
            sizes.get(ARC_CONTENT_TYPES, 0) > MAX_XLSX_CONTENT_TYPES_BYTES
            or sizes.get(ARC_STYLE, 0) > MAX_XLSX_STYLES_BYTES
        ):
            raise InvalidInputError(WORKBOOK_TOO_LARGE_MESSAGE)
        try:
            shared_strings = _shared_strings_part(archive)
        except _WORKBOOK_ERRORS as exc:
            raise InvalidInputError(WORKBOOK_UNREADABLE_MESSAGE) from exc
        if (
            shared_strings is not None
            and sizes.get(shared_strings, 0) > MAX_XLSX_SHARED_STRINGS_BYTES
        ):
            raise InvalidInputError(WORKBOOK_TOO_LARGE_MESSAGE)


def read_xlsx(data: bytes, *, check_token: str) -> ImportSheet:
    """The first worksheet's stored values, columns A–BL."""
    reject_empty(data)
    if not data.startswith(_ZIP_SIGNATURE):
        raise InvalidInputError(NOT_XLSX_MESSAGE)
    _check_zip(data)
    records: list[ImportRecord] = []
    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(data), read_only=True, data_only=True, keep_links=False
        )
        try:
            worksheets = workbook.worksheets
            if not worksheets or not isinstance(worksheets[0], ReadOnlyWorksheet):
                raise InvalidInputError(WORKBOOK_UNREADABLE_MESSAGE)
            worksheet = worksheets[0]
            title = worksheet.title
            if worksheet.sheet_state != "visible":
                raise InvalidInputError(
                    f"The first worksheet ({title}) is hidden. PartFlow reads the first worksheet"
                    " only — move the sheet with the import rows to the first position and make"
                    " it visible, then check the file again."
                )
            worksheet.reset_dimensions()
            rows = worksheet.iter_rows(
                min_row=1, min_col=1, max_col=MAX_IMPORT_COLUMNS, values_only=False
            )
            filler_row: object = None
            for index, cells in enumerate(rows):
                if cells is filler_row:
                    continue
                row = 1 + index
                values = tuple(_cell_value(cell, row) for cell in cells)
                if all(is_blank_cell(value) for value in values):
                    if all(cell is _FILLER for cell in cells):
                        # openpyxl yields one shared filler tuple for every
                        # absent row: skip the rest of a gap without work.
                        filler_row = cells
                    continue
                records.append(ImportRecord(row=row, cells=values))
                if len(records) >= MAX_IMPORT_RECORDS:
                    break
        finally:
            workbook.close()
    except (*_WORKBOOK_ERRORS, _UnreadableCellError) as exc:
        raise InvalidInputError(WORKBOOK_UNREADABLE_MESSAGE) from exc
    return ImportSheet(
        file_format=ImportFormat.XLSX,
        records=tuple(records),
        check_token=check_token,
        worksheet=title,
    )


# ---------------------------------------------------------------------------
# Templates
# ---------------------------------------------------------------------------

#: Number format and width per template column (OD-S1-5): the identity
#: columns are Text so leading zeros stay.
_TEMPLATE_LAYOUT: Final = {
    COLUMN_WORK_ORDER_NUMBER: ("@", 22),
    COLUMN_PART_NUMBER: ("@", 26),
    COLUMN_REQUESTED_QUANTITY: ("0", 20),
    COLUMN_JOB_NUMBER: ("@", 18),
    COLUMN_DUE_DATE: ("yyyy-mm-dd", 14),
}


def write_xlsx_template(columns: tuple[str, ...]) -> bytes:
    """One sheet ``Work Orders``: the bold, frozen header row only."""
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Work Orders"
    bold = Font(bold=True)
    for index, column in enumerate(columns, start=1):
        cell = worksheet.cell(row=1, column=index, value=column)
        cell.font = bold
        number_format, width = _TEMPLATE_LAYOUT.get(column, ("General", 16))
        letter = cell.column_letter
        worksheet.column_dimensions[letter].number_format = number_format
        worksheet.column_dimensions[letter].width = width
    worksheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
